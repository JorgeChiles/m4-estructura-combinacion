#!/usr/bin/env python3
"""
OWA oficial de M4 para los 26 modelos de la tesis.

Que hace
--------
Los consolidados guardan, en cada hoja, los PRONOSTICOS de cada modelo
(bloque que arranca en la fila "Datos Prueba"). Con eso se pueden calcular
las metricas oficiales de la competencia y poner los 26 modelos en la misma
escala con la que se midieron los 61 metodos de M4.

    sMAPE = (200/h) * sum |A - F| / (|A| + |F|)
    MASE  = mean|A - F| / naive estacional in-sample
    OWA   = 0.5 * (sMAPE/sMAPE_Naive2 + MASE/MASE_Naive2)

Por que NO se usa el SMAPE ya guardado
--------------------------------------
El del notebook es 100*mean(2|F-A|/(|A|+|F|+1e-6)), que es la misma formula
que la de M4 salvo el epsilon. Pero el MASE no esta guardado y hace falta el
pronostico igual, asi que se recalcula todo desde los pronosticos: una sola
implementacion, la del script 06, para modelos y benchmarks por igual.

Fuente de verdad
----------------
Los valores reales y las series de entrenamiento se leen de los CSV
ORIGINALES de M4 (m4-structural-complexity/data/), no de los transpuestos.
Asi el denominador del MASE y el Naive2 son identicos a los del script 06 y
los OWA son comparables. El script ademas verifica que los reales guardados
en el consolidado coincidan con los oficiales.

Salidas
-------
    owa_modelos_por_serie.csv    una fila por (serie, modelo)
    owa_modelos_resumen.xlsx     OWA por modelo y frecuencia
    owa_modelos_ranking.xlsx     ranking global

Uso:  python3 owa_modelos.py
"""
import glob
import os
import re
import sys
import time
import warnings
from pathlib import Path

os.environ.setdefault("PYTHONWARNINGS", "ignore")

import numpy as np
import pandas as pd

warnings.filterwarnings("ignore")

AQUI = Path(__file__).resolve().parent
COMPLEJIDAD = AQUI.parent / "m4-structural-complexity"

# Se reutilizan las funciones del script 06 en vez de reimplementarlas: si
# se corrige una formula alla, esto la hereda.
sys.path.insert(0, str(COMPLEJIDAD / "src"))
import importlib
M4 = importlib.import_module("06_m4_metrics")

DATA = COMPLEJIDAD / "data"


# =========================================================
# CONFIGURACION
# =========================================================

# carpeta, prefijo del consolidado, nombre M4, h, m
FRECS = [
    ("YEARLY_DATA",   "resultados_consolidados_yearly",    "Yearly",     6,  1),
    ("QUARTELY_DATA", "resultados_consolidados_quarterly", "Quarterly",  8,  4),
    ("MONTHLY_DATA",  "resultados_consolidados",           "Monthly",   18, 12),
    ("WEEKLY_DATA",   "resultados_consolidados_weekly",    "Weekly",    13,  1),
    ("DAILY_DATA",    "resultados_consolidados",           "Daily",     14,  1),
    ("HOURLY_DATA",   "resultados_consolidados_hourly",    "Hourly",    48, 24),
]

# copias de respaldo que no son resultados vigentes
EXCLUIR = ("ORIGINAL", "PREVIO", "ANTES", "LIMPIO", "RECORRIDA", "PRE_", "CORRUPTO")

BENCHMARKS = ("naive2", "naive", "snaive")


# =========================================================
# LECTURA DE LOS CONSOLIDADOS
# =========================================================

# El notebook, cuando un modelo YA tenia pronostico para esa serie, agrega
# la columna con sufijo en vez de pisarla:
#     while col in df_pron.columns: col = f"{base} v{k}"
# O sea que "LSTM v2" no es otro modelo: es la MISMA serie recorrida dos
# veces. Se conserva la primera y se descartan las repeticiones, para que
# haya exactamente un pronostico por (serie, modelo).
RE_RECORRIDA = re.compile(r"\s+v\d+$")


def pronosticos_de_hoja(filas):
    """Del bloque que arranca en 'Datos Prueba' devuelve
    ({modelo: array}, n_repeticiones_descartadas).
    Devuelve (None, 0) si la hoja no tiene bloque de pronosticos."""
    i_cab = next((i for i, f in enumerate(filas)
                  if f and str(f[0]).strip() == "Datos Prueba"), None)
    if i_cab is None:
        return None, 0

    cab = filas[i_cab]
    cols, vistos, repetidas = {}, set(), 0
    for j, v in enumerate(cab):
        if v is None or str(v).strip() in ("", "nan"):
            continue
        nombre = RE_RECORRIDA.sub("", str(v).strip())
        if nombre in vistos:
            repetidas += 1
            continue
        vistos.add(nombre)
        cols[j] = nombre

    datos = {j: [] for j in cols}
    for f in filas[i_cab + 1:]:
        if not f or f[0] is None:
            break
        for j in cols:
            v = f[j] if j < len(f) else None
            datos[j].append(np.nan if v is None else
                            (float(v) if isinstance(v, (int, float)) else np.nan))

    return {cols[j]: np.asarray(v, float) for j, v in datos.items()}, repetidas


def leer_consolidados(carpeta, base):
    """Todas las hojas de todos los consolidados de la carpeta.
    Si una serie aparece en varios archivos se toma la primera: son la misma
    corrida, los duplicados ya se limpiaron en su momento."""
    import openpyxl
    archivos = [a for a in sorted(glob.glob(f"{carpeta}/{base}*.xlsx"))
                if not any(k in a for k in EXCLUIR)]
    series, repetidas = {}, 0
    for arch in archivos:
        try:
            wb = openpyxl.load_workbook(arch, read_only=True, data_only=True)
        except Exception as e:
            print(f"   ⚠ no se pudo abrir {os.path.basename(arch)}: {e}", flush=True)
            continue
        for hoja in wb.sheetnames:
            if hoja in series:
                continue
            filas = [f for f in wb[hoja].iter_rows(values_only=True)]
            pron, rep = pronosticos_de_hoja(filas)
            repetidas += rep
            if pron:
                series[hoja] = pron
        wb.close()
    return series, repetidas


# =========================================================
# CALCULO
# =========================================================

def procesar(carpeta, base, nombre, h, m):
    print(f"── {nombre}  (h={h}, m={m})", flush=True)
    t0 = time.time()

    pron_por_serie, repetidas = leer_consolidados(carpeta, base)
    print(f"   {len(pron_por_serie)} series leidas  ({time.time()-t0:.0f} s)", flush=True)
    if repetidas:
        print(f"   {repetidas} pronosticos repetidos (serie recorrida mas de una "
              f"vez); se conserva el primero", flush=True)

    tr = M4.load_m4(DATA / f"{nombre}-train.csv")
    te = M4.load_m4(DATA / f"{nombre}-test.csv")

    filas, desajustes, sin_datos = [], 0, 0

    for sid, pron in sorted(pron_por_serie.items()):
        if sid not in tr or sid not in te:
            sin_datos += 1
            continue
        entren = np.asarray(tr[sid], float)
        real = np.asarray(te[sid], float)[:h]

        # control: los reales guardados en el consolidado tienen que ser los
        # oficiales. Si no coinciden, la evaluacion de esa serie no es valida.
        guardado = pron.get("Datos Prueba")
        if guardado is not None and len(guardado) >= len(real):
            if not np.allclose(guardado[:len(real)], real, rtol=1e-3, atol=1e-3,
                               equal_nan=True):
                desajustes += 1

        # benchmarks, con las MISMAS funciones que el script 06
        bench = {
            "naive2": M4.forecast_naive2(entren, h, m),
            "naive":  M4.forecast_naive(entren, h),
            "snaive": M4.forecast_snaive(entren, h, m),
        }

        for modelo, pred in list(pron.items()) + list(bench.items()):
            if modelo == "Datos Prueba":
                continue
            pred = np.asarray(pred, float)
            n = min(len(pred), len(real))
            if n == 0 or not np.all(np.isfinite(pred[:n])):
                filas.append(dict(frecuencia=nombre, serie=sid, modelo=modelo,
                                  smape=np.nan, mase=np.nan, n_pron=n,
                                  completo=False))
                continue
            filas.append(dict(
                frecuencia=nombre, serie=sid, modelo=modelo,
                smape=M4.smape(real[:n], pred[:n]),
                mase=M4.mase(entren, real[:n], pred[:n], m),
                n_pron=n, completo=(n == h)))

    if desajustes:
        print(f"   ⚠ {desajustes} series con datos de prueba distintos de los "
              f"oficiales de M4", flush=True)
    if sin_datos:
        print(f"   ⚠ {sin_datos} series del consolidado no estan en M4", flush=True)
    return pd.DataFrame(filas)


def resumir(df):
    """OWA por modelo y frecuencia. Como en M4, sobre PROMEDIOS.

    El denominador Naive2 se restringe a LAS MISMAS series que pudo
    pronosticar el modelo. Sin esto, Holt-Winters -- que solo existe donde
    m>1 (2.414 de 4.773 series) -- quedaria dividido por un Naive2
    promediado sobre series que el modelo nunca vio, y su OWA no seria
    comparable con el del resto."""
    n2 = (df[df.modelo == "naive2"]
          .drop_duplicates(["frecuencia", "serie"])
          .set_index(["frecuencia", "serie"])[["smape", "mase"]])

    filas = []
    grupos = [(f, g) for f, g in df.groupby("frecuencia")] + [("TODAS", df)]
    for frec, g in grupos:
        for modelo, gm in g.groupby("modelo"):
            val = gm.dropna(subset=["smape", "mase"])
            if val.empty:
                continue
            base = n2.reindex(pd.MultiIndex.from_frame(val[["frecuencia", "serie"]]))
            base_s, base_m = base.smape.mean(), base.mase.mean()
            s, ma = val.smape.mean(), val.mase.mean()
            filas.append({
                "frecuencia": frec, "modelo": modelo,
                "n_series": len(gm), "n_validas": len(val),
                "cobertura_%": round(100 * len(val) / len(gm), 1),
                "sMAPE": s, "MASE": ma,
                "OWA": 0.5 * (s / base_s + ma / base_m) if base_s and base_m else np.nan,
                "sMAPE_mediana": val.smape.median(),
                "MASE_mediana": val.mase.median(),
                # OWA con medianas: diagnostico para ver si el promedio esta
                # dominado por unas pocas series que divergen. NO es la
                # metrica oficial ni es comparable con los OWA publicados.
                "OWA_mediana": 0.5 * (val.smape.median() / base.smape.median()
                                      + val.mase.median() / base.mase.median()),
                "mase_extremos": int((val.mase > 100).sum()),
            })
    return pd.DataFrame(filas)


# =========================================================
# OWA REPONDERADO A LA POBLACION DE M4
# =========================================================

# Series de cada frecuencia en M4 completo. La muestra de la tesis NO tiene
# estas proporciones: Daily entra con 1000 de 4227 y Hourly con 414 de 414,
# asi que pesan 21% y 8.7% en vez de 4.2% y 0.4%. Como son las frecuencias
# donde varios modelos se descontrolan, el promedio simple sobre la muestra
# los castiga mucho mas de lo que los castigaria M4.
POBLACION_M4 = {"Yearly": 23000, "Quarterly": 24000, "Monthly": 48000,
                "Weekly": 359, "Daily": 4227, "Hourly": 414}


def ponderado(res):
    """OWA global con los pesos poblacionales de M4.

    Se ponderan sMAPE y MASE por separado y recien despues se hace el
    cociente, que es como agrega la competencia (razon de promedios, no
    promedio de razones)."""
    tot = sum(POBLACION_M4.values())
    porf = res[res.frecuencia != "TODAS"]
    filas = []
    for modelo, g in porf.groupby("modelo"):
        g = g[g.frecuencia.isin(POBLACION_M4)]
        w = g.frecuencia.map(POBLACION_M4) / tot
        # si al modelo le falta alguna frecuencia (Holt-Winters solo existe
        # con m>1), se renormaliza sobre las que si tiene
        w = w / w.sum()
        base = porf[(porf.modelo == "naive2") & (porf.frecuencia.isin(g.frecuencia))]
        wb = base.frecuencia.map(POBLACION_M4) / tot
        wb = wb / wb.sum()
        s, ma = (g.sMAPE * w).sum(), (g.MASE * w).sum()
        bs, bm = (base.sMAPE * wb).sum(), (base.MASE * wb).sum()
        filas.append({"modelo": modelo, "frecuencias": len(g),
                      "sMAPE_pond": s, "MASE_pond": ma,
                      "OWA_pond": 0.5 * (s / bs + ma / bm) if bs and bm else np.nan})
    return pd.DataFrame(filas).sort_values("OWA_pond").reset_index(drop=True)


# =========================================================
# PRINCIPAL
# =========================================================

def main():
    t0 = time.time()
    partes = []
    for carpeta, base, nombre, h, m in FRECS:
        if not os.path.isdir(AQUI / carpeta):
            print(f"── {nombre}: falta la carpeta, se omite", flush=True)
            continue
        partes.append(procesar(str(AQUI / carpeta), base, nombre, h, m))

    df = pd.concat(partes, ignore_index=True)
    df.to_csv(AQUI / "owa_modelos_por_serie.csv", index=False)

    res = resumir(df)
    res.to_excel(AQUI / "owa_modelos_resumen.xlsx", index=False)

    glob_ = (res[res.frecuencia == "TODAS"]
             .sort_values("OWA")[["modelo", "n_validas", "cobertura_%", "sMAPE",
                                  "MASE", "OWA", "OWA_mediana", "mase_extremos"]]
             .reset_index(drop=True))
    glob_.insert(0, "puesto", range(1, len(glob_) + 1))
    glob_.to_excel(AQUI / "owa_modelos_ranking.xlsx", index=False)

    print(f"\n═══ RANKING GLOBAL POR OWA ({(time.time()-t0)/60:.1f} min, "
          f"{df.serie.nunique()} series) ═══")
    print(glob_.round(3).to_string(index=False))

    print("\n═══ OWA POR FRECUENCIA ═══")
    piv = res[res.frecuencia != "TODAS"].pivot(index="modelo",
                                               columns="frecuencia", values="OWA")
    orden = [f[2] for f in FRECS if f[2] in piv.columns]
    print(piv[orden].round(3).sort_values(orden[0]).to_string())

    pond = ponderado(res)
    pond.to_excel(AQUI / "owa_modelos_ponderado.xlsx", index=False)
    print("\n═══ OWA REPONDERADO A LA POBLACION DE M4 ═══")
    print("(la muestra sobre-representa Daily y Hourly; esto corrige la mezcla)")
    print(pond.round(3).to_string(index=False))

    ext = res[(res.frecuencia == "TODAS") & (res.mase_extremos > 0)]
    if len(ext):
        print("\n⚠ modelos con MASE extremos (>100), que inflan su promedio:")
        print(ext[["modelo", "mase_extremos", "MASE", "MASE_mediana"]]
              .sort_values("mase_extremos", ascending=False).round(3).to_string(index=False))

    print("\nowa_modelos_por_serie.csv\nowa_modelos_resumen.xlsx\nowa_modelos_ranking.xlsx")
    return 0


if __name__ == "__main__":
    sys.exit(main())
